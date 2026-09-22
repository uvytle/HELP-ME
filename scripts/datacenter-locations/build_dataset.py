#!/usr/bin/env python3
"""
Scrape US data center locations from a few open sources and write a simple
3-column spreadsheet: Data Center Name, Location, Type.

Install:
    pip install -r requirements.txt

Run:
    python build_dataset.py

Sources (see README.md for why each one was picked, and why datacenters.com /
datacentermap.com are NOT included — both block automated access):
    - usdatamap.com   (public static dataset)
    - PeeringDB       (public API, colocation/carrier-neutral facilities)
    - OpenStreetMap   (Overpass API, telecom=data_center)
    - cloud_regions   (hand-maintained list of AWS/Azure/GCP/OCI/IBM US regions)

Output: data/datacenters_us.xlsx
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from sources import cloud_regions, osm_overpass, peeringdb, usdatamap

OUTPUT_PATH = Path(__file__).parent / "data" / "datacenters_us.xlsx"

SOURCE_FETCHERS = {
    "usdatamap": usdatamap.fetch,
    "peeringdb": peeringdb.fetch,
    "osm": osm_overpass.fetch,
    "cloud_regions": cloud_regions.fetch,
}


def collect_records() -> list[dict]:
    records = []
    for name, fetcher in SOURCE_FETCHERS.items():
        try:
            recs = fetcher()
            print(f"  {name}: {len(recs)} records")
        except Exception as exc:  # noqa: BLE001 - one source failing shouldn't kill the run
            print(f"  {name}: FAILED ({exc})")
            recs = []
        records.extend(recs)
    return records


def location_string(r: dict) -> str:
    if r.get("city") and r.get("state_code"):
        return f"{r['city']}, {r['state_code']}"
    if r.get("address"):
        return r["address"]
    return r.get("state") or ""


def write_excel(records: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Centers"

    ws.append(["Data Center Name", "Location", "Type"])
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)

    for r in records:
        ws.append([r.get("name"), location_string(r), r.get("type")])
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = Font(name="Arial")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(records) + 1}"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} ({len(records):,} rows)")


def main() -> None:
    print("Fetching sources...")
    records = collect_records()
    write_excel(records)


if __name__ == "__main__":
    main()
